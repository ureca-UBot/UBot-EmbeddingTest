"""Validate labels and source references without importing an embedding model."""
from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter
from pathlib import Path

from jsonschema import Draft202012Validator
import openpyxl

from retrieval_common import semantic_contract_errors

ROOT = Path(__file__).resolve().parents[1]


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def reference_catalog(audit, source_paths):
    verified = []
    current = None
    for source in audit["sources"]:
        path = Path(source_paths.get(source["key"], source["path"]))
        with path.open("rb") as stream:
            digest = hashlib.file_digest(stream, "sha256").hexdigest()
        if digest != source["sha256"]:
            raise ValueError(f"source SHA256 changed: {source['key']}")
        verified.append(source["key"])
        if source["key"] == "current_retrieval_base":
            current = path
    if current is None:
        raise ValueError("current_retrieval_base missing from source audit")
    wb = openpyxl.load_workbook(current, read_only=True, data_only=True)
    try:
        base = {r[0]: {"question": r[2], "answer": r[3]} for r in wb["FAQ 원문"].iter_rows(min_row=2, values_only=True) if r[0]}
        catalog = {"current-faq-1024-v1": base}
        fixture_groups = {}
        for r in wb["시험 문서"].iter_rows(min_row=2, values_only=True):
            if not r[1]:
                continue
            fixture_groups.setdefault(r[0], {})[r[1]] = {"base_faq_id": r[2], "question": r[3], "answer": r[4], "valid_from": r[6], "valid_until": r[7]}
        # Source contracts for the design examples. Never merge TEST docs into base.
        cf = fixture_groups["CF-0061"]
        catalog["fixture-CF-0061-v1"] = {**{f"FAQ-{i:03d}": base[f"FAQ-{i:03d}"] for i in range(18, 28)}, **cf}
        ad = fixture_groups["AD-0141"]
        replaced = {doc["base_faq_id"] for doc in ad.values()}
        catalog["fixture-AD-0141-v1"] = {**{fid: doc for fid, doc in base.items() if fid not in replaced}, **ad}
        return catalog, verified
    finally:
        wb.close()


def validate(cases, schema, catalog, allow_design_examples=False):
    Draft202012Validator.check_schema(schema)
    validator = Draft202012Validator(schema)
    errors = []
    for row, case in enumerate(cases, 1):
        for error in validator.iter_errors(case):
            errors.append(f"row {row} {case.get('case_id', '?')} {list(error.path)}: {error.message}")
    if errors:
        return errors
    for case in cases:
        status = case["annotation"]["status"]
        if status != "approved" and not (allow_design_examples and status == "design_example" and case["split"] == "design_only"):
            errors.append(f"{case['case_id']}: unapproved annotation cannot enter evaluation")
    errors.extend(semantic_contract_errors(cases, catalog))
    if not cases:
        errors.append("dataset is empty")
    return errors


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--cases", type=Path, default=ROOT / "faq/retrieval_design_examples.jsonl")
    parser.add_argument("--schema", type=Path, default=ROOT / "schemas/retrieval_case.schema.json")
    parser.add_argument("--audit", type=Path, default=ROOT / "configs/retrieval_source_audit.json")
    parser.add_argument("--source-map", type=Path)
    parser.add_argument("--allow-design-examples", action="store_true")
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    try:
        cases = [json.loads(line) for line in args.cases.read_text(encoding="utf-8").splitlines() if line.strip()]
        schema = read_json(args.schema)
        catalog, verified = reference_catalog(read_json(args.audit), read_json(args.source_map) if args.source_map else {})
        errors = validate(cases, schema, catalog, args.allow_design_examples)
        result = {"valid": not errors, "case_count": len(cases), "corpus_faq_count": len(catalog["current-faq-1024-v1"]),
                  "source_hashes_verified": verified, "cohorts": dict(Counter(c["reporting_cohort"] for c in cases)) if not errors else None,
                  "schema_draft": "2020-12", "design_examples_allowed": args.allow_design_examples,
                  "models_executed": False, "errors": errors}
    except (OSError, ValueError, KeyError) as exc:
        result = {"valid": False, "models_executed": False, "errors": [str(exc)]}
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["valid"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
