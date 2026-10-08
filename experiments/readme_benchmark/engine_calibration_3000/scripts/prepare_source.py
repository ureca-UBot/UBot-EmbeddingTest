"""Freeze exactly the user's selected 3,000 rows; never execute workbook prose."""
import hashlib,json,re
from collections import Counter
from pathlib import Path
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
SOURCE=Path(r'C:\Users\eongp\Downloads\FAQ_RAG_15개항목_각200건_총3000건_피드백수정본(1).xlsx')

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def write_lines(p,rows):p.write_text(''.join(json.dumps(r,ensure_ascii=False)+'\n' for r in rows),encoding='utf-8')
def main():
    out=ROOT/'data';out.mkdir(parents=True,exist_ok=True)
    if (out/'manifest.json').exists():raise SystemExit('Source already frozen; do not overwrite')
    wb=openpyxl.load_workbook(SOURCE,read_only=True,data_only=True)
    it=wb['테스트 3000건'].iter_rows(values_only=True);headers=next(it)
    raw=[dict(zip(headers,r)) for r in it if r[0]]
    assert len(raw)==3000 and all(n==200 for n in Counter(r['항목 코드'] for r in raw).values())
    docs=[]
    for r in wb['FAQ 원문'].iter_rows(min_row=2,values_only=True):
        if r[0]:docs.append({'faq_id':r[0],'category':r[1],'question':r[2],'answer':r[3],'intent':r[4]})
    wb.close();assert len(docs)==1024
    ids={d['faq_id'] for d in docs};cases=[]
    cohorts={'SF':'general','NC':'general','PS':'general','CE':'condition','MC':'condition',
             'PI':'partial','MT':'contextual','UI':'personal','CF':'conflict','AR':'api','AD':'adversarial',
             'SR':'no_faq','HR':'no_faq','EC':'no_faq','RT':'repeat'}
    for rownum,r in enumerate(raw,2):
        q=r['사용자 질문'];assert isinstance(q,str) and q
        history=json.loads(r['대화 이력 (JSON)']) if r['대화 이력 (JSON)'] else []
        assert isinstance(history,list)
        model_query=('대화 이력:\n'+'\n'.join(h['role']+': '+h['content'] for h in history)+'\n현재 질문: '+q) if history else q
        gold=list(dict.fromkeys(re.findall(r'FAQ-\d+',r['근거 원문 FAQ ID (추적용)'] or '')))
        assert set(gold)<=ids,(r['실행 ID'],gold)
        kind=r['항목 코드'];group=r['독립 집계 단위'] or r['실행 ID']
        # Generation ABSTAIN alone is never converted to global NO_FAQ.
        nofaq=kind in ['SR','HR','EC'] and not gold and group.startswith('NO_ANSWER_TOPIC:')
        cohort=cohorts[kind]
        if kind in ['SR','HR','EC'] and not nofaq:cohort='context_abstention'
        cases.append({'case_id':r['실행 ID'],'source_row':rownum,'type':kind,'type_name':r['테스트 항목'],
                      'query':q,'history':history,'model_query':model_query,'cohort':cohort,
                      'family_id':group,'source_ids':gold,'no_faq_truth':nofaq,
                      'fact_groups':[{'fact_id':'SOURCE:'+i,'primary_ids':[i],'acceptable_ids':[]} for i in gold],
                      'expected_generation_status':r['기대 상태'],'expected_route':r['기대 처리 경로'],
                      'repeat_index':r['실행 회차'],'repeat_total':r['총 반복 횟수'],
                      'annotation':'workbook tracking IDs; source coverage proxy, not independently reviewed atomic facts'})
    assert len({c['case_id'] for c in cases})==3000
    assert [c['query'] for c in cases]==[r['사용자 질문'] for r in raw]
    write_lines(out/'source_rows.jsonl',raw);write_lines(out/'cases.jsonl',cases);write_lines(out/'corpus.jsonl',docs)
    manifest={'source_workbook':str(SOURCE),'source_sha256':sha(SOURCE),'question_sheet':'테스트 3000건',
              'rows':3000,'corpus_rows':1024,'type_counts':dict(Counter(c['type'] for c in cases)),
              'cohort_counts':dict(Counter(c['cohort'] for c in cases)),
              'unique_raw_questions':len({c['query'] for c in cases}),
              'query_text_changed':False,'rows_dropped':0,'rows_added':0,'holdout_split':False,
              'history_rule':'unchanged current question with fixed original history concatenation when present',
              'provided_context_used_for_retrieval':False,'answers_and_gold_used_as_model_input':False,
              'original_api_and_persona_metadata_retained_in_source_rows_only':True,
              'later_user_authorization':'use this selected 3000-question workbook unchanged for each engine',
              'hashes':{p.name:sha(p) for p in [out/'source_rows.jsonl',out/'cases.jsonl',out/'corpus.jsonl']}}
    (out/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(manifest,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
