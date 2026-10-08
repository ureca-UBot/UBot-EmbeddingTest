"""Read back the final workbook and compare every question to the frozen JSONL."""
import json
import xml.etree.ElementTree as ET
from zipfile import ZipFile
from pathlib import Path
import openpyxl
from build_variations import read_jsonl,write_json,sha

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs'
path=OUT/'FAQ_3000쌍_각10개_변형질문_30000행.xlsx'
variants=read_jsonl(ROOT/'data'/'variations.jsonl')
source=read_jsonl(ROOT/'data'/'source_pairs.jsonl')
manifest=json.loads((ROOT/'data'/'manifest.json').read_text(encoding='utf-8'))
catalog=json.loads((ROOT/'data'/'purpose_catalog.json').read_text(encoding='utf-8'))
wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
assert wb.sheetnames==['변형30000','검증목적','원본3000']
sheet=wb['변형30000']
assert sheet['B4'].value==3000 and sheet['F4'].value==30000
main=list(sheet.iter_rows(min_row=9,values_only=True))
assert len(main)==30000,len(main)
for excel,row in zip(main,variants):
    assert excel[:6]==(row['parent_case_id'],row['case_id'],row['type'],row['primary_variation'],row['query'],row['expected_answer']),row['case_id']
original=list(wb['원본3000'].iter_rows(min_row=5,values_only=True))
assert len(original)==3000
for excel,row in zip(original,source):
    assert excel[:4]==(row['실행 ID'],row['항목 코드'],row['사용자 질문'],row['정답 예시']),row['실행 ID']
counts=list(wb['검증목적'].iter_rows(min_row=6,max_row=21,values_only=True))
for excel,entry in zip(counts,catalog):
    tag=entry['tag']
    assert excel[0]==tag
    assert excel[4]==manifest['primary_variation_counts'].get(tag,0),(tag,excel[4])
    assert excel[5]==manifest['validation_tag_counts'].get(tag,0),(tag,excel[5])
wb.close()
namespace='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
with ZipFile(path) as archive:
    frozen={}
    for filename in ['xl/worksheets/sheet1.xml','xl/worksheets/sheet3.xml']:
        with archive.open(filename) as stream:
            for event,element in ET.iterparse(stream,events=['start']):
                if element.tag==namespace+'pane':frozen[filename]=element.attrib
                if element.tag==namespace+'sheetData':break
    assert frozen['xl/worksheets/sheet1.xml']['topLeftCell']=='C9'
    assert frozen['xl/worksheets/sheet3.xml']['topLeftCell']=='C5'
    tables=[ET.fromstring(archive.read(name)).attrib for name in archive.namelist() if name.startswith('xl/tables/') and name.endswith('.xml')]
    assert {t['ref'] for t in tables}=={'A8:P30008','A4:I3004'}
code_and_data={}
for folder in ['data','scripts','outputs']:
    for f in (ROOT/folder).iterdir():
        if f.is_file() and f.name not in ['export_verification.json','artifact_inventory.json']:
            code_and_data[str(f.relative_to(ROOT))]={'sha256':sha(f),'bytes':f.stat().st_size}
for name in ['README.md','compose.yaml','.gitignore']:
    code_and_data[name]={'sha256':sha(ROOT/name),'bytes':(ROOT/name).stat().st_size}
report={'passed':True,'exported_variation_rows':len(main),'exported_original_rows':len(original),
        'all_question_answer_cells_match_jsonl':True,'all_purpose_counts_match_manifest':True,
        'header_and_identifier_freeze_preserved':True,'table_ranges_correct':True,
        'workbook_sha256':sha(path),'independent_semantic_review_complete':False,'benchmark_executed':False}
write_json(OUT/'export_verification.json',report)
write_json(OUT/'artifact_inventory.json',code_and_data)
print(json.dumps(report,ensure_ascii=False,indent=2))
